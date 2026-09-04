from __future__ import annotations

import json
import zipfile
from datetime import date
from io import BytesIO

import numpy as np
import pandas as pd
import pytest
from shapely.geometry import Polygon

from core.crs import EPSG_32638_CRS_MODE, local_crs, preset_crs_32638
from core.engineering_controls import REGION_CONTROL_SOURCE, create_control_point, create_control_region
from core.project_io import load_project_archive, save_project_archive
from core.scenarios import create_map_scenario, scenario_to_generated_map
from utils.export import (
    GEOTIFF_NODATA,
    ZMAP_NULL_VALUE,
    batch_map_package_zip_bytes,
    build_map_export_metadata,
    control_regions_export_dataframe,
    engineering_controls_export_dataframe,
    grid_to_excel_bytes,
    grid_to_geotiff_bytes,
    grid_to_zmap_ascii_bytes,
    map_geotiff_export_files,
    map_zmap_export_files,
    parse_zmap_grid_ascii_bytes,
)


def _asymmetric_grid():
    x_axis = np.array([428000.0, 428250.0, 428500.0], dtype=float)
    y_axis = np.array([3365000.0, 3365250.0, 3365500.0], dtype=float)
    grid_x, grid_y = np.meshgrid(x_axis, y_axis)
    grid_z = np.array(
        [
            [1.0, np.nan, 3.0],
            [4.0, 5.0, 6.0],
            [7.0, 8.0, 99.0],
        ],
        dtype=float,
    )
    return grid_x, grid_y, grid_z


def _map_result(*, crs=None, layer="Upper", variance=False):
    grid_x, grid_y, grid_z = _asymmetric_grid()
    grid_variance = None
    if variance:
        grid_variance = np.array(
            [
                [1.0, 4.0, 9.0],
                [16.0, 25.0, 36.0],
                [49.0, 64.0, 81.0],
            ],
            dtype=float,
        )
    control = create_control_point(
        x=428250.0,
        y=3365250.0,
        property_name="Pressure",
        value=3120.0,
        property_unit="psi",
        reservoir_layer=layer,
        panel="A",
        pressure_reference_date=date(2026, 1, 1),
        control_id="CP-001",
    )
    region = create_control_region(
        region_name="Support",
        geometry=Polygon(
            [
                (428000.0, 3365000.0),
                (428500.0, 3365000.0),
                (428500.0, 3365500.0),
                (428000.0, 3365500.0),
            ]
        ),
        property_name="Pressure",
        target_value=3150.0,
        property_unit="psi",
        reservoir_layer=layer,
        panel="A",
        pressure_reference_date="2026-01-01",
        control_point_spacing=250.0,
        region_id="CR-001",
    )
    controls = pd.DataFrame(
        [
            {
                "Control_ID": "CP-001",
                "Source_Type": "Manual Control Point",
                "X": 428250.0,
                "Y": 3365250.0,
                "Pressure": 3120.0,
                "Property_Unit": "psi",
                "Reservoir_Layer": layer,
                "Panel": "A",
                "Pressure_Map_Reference_Date": "2026-01-01",
                "Active": True,
                "Comment": "manual support",
            },
            {
                "Control_ID": "CR-001-GC-001",
                "Source_Type": REGION_CONTROL_SOURCE,
                "Region_ID": "CR-001",
                "X": 428500.0,
                "Y": 3365500.0,
                "Pressure": 3150.0,
                "Property_Unit": "psi",
                "Reservoir_Layer": layer,
                "Panel": "A",
                "Pressure_Map_Reference_Date": "2026-01-01",
                "Active": True,
                "Comment": "region support",
            },
        ]
    )
    return {
        "grid_x": grid_x,
        "grid_y": grid_y,
        "grid_z": grid_z,
        "grid_variance": grid_variance,
        "panel_grid": np.array([["A", None, "A"], ["A", "A", "A"], ["A", "A", "A"]], dtype=object),
        "property_col": "Pressure",
        "property_type": "Pressure",
        "unit": "psi",
        "coordinate_unit": "meters",
        "crs": crs or preset_crs_32638().to_dict(),
        "is_pressure_map": True,
        "map_reference_date": date(2026, 1, 1),
        "x_col": "X",
        "y_col": "Y",
        "method": "Ordinary Kriging" if variance else "IDW",
        "method_parameters": {
            "variogram_model": "Spherical",
            "variogram_range_convention": "Practical Range",
            "range": 1000.0,
            "variance": 200.0,
            "nugget": 10.0,
        }
        if variance
        else {"power": 2.0, "neighbors": 12},
        "grid_parameters": {"nx": 3, "ny": 3, "buffer_fraction": 0.0},
        "mask_parameters": {"mode": "Reservoir Boundary"},
        "mask_info": {
            "mask_mode": "Reservoir Boundary",
            "valid_grid_cells": int(np.isfinite(grid_z).sum()),
            "grid_cells_before_mask": int(grid_z.size),
            "masked_cells": int(np.isnan(grid_z).sum()),
        },
        "duplicate_method": "Average",
        "selected_panels": ["A"],
        "selected_layers": [layer],
        "reservoir_layer": layer,
        "layer_mapping_scope": "All Layers",
        "panel_interpolation_mode": "Combined Selected Panels",
        "geometry_references": {"reservoir_boundary_name": "R1 Boundary", "selected_panel_names": ["A"]},
        "included_observations": pd.DataFrame(),
        "excluded_observations": pd.DataFrame(),
        "engineering_controls": controls,
        "engineering_control_points": [control],
        "engineering_control_regions": [region],
        "engineering_control_count": 2,
        "control_region_count": 1,
        "measured_observation_count": 5,
        "title": f"Pressure {layer} 2026",
        "export_metadata": {
            "Property": "Pressure",
            "Property_Unit": "psi",
            "Reservoir_Layer": layer,
            "Pressure_Map_Reference_Date": "2026-01-01",
        },
    }


def test_geotiff_round_trip_epsg32638_orientation_bounds_and_nodata():
    rasterio = pytest.importorskip("rasterio")
    grid_x, grid_y, grid_z = _asymmetric_grid()
    data = grid_to_geotiff_bytes(grid_x, grid_y, grid_z, preset_crs_32638().to_dict())

    with rasterio.io.MemoryFile(data) as memory_file:
        with memory_file.open() as dataset:
            values = dataset.read(1)
            assert dataset.width == 3
            assert dataset.height == 3
            assert dataset.crs.to_epsg() == 32638
            assert dataset.nodata == GEOTIFF_NODATA
            assert dataset.transform.a == pytest.approx(250.0)
            assert dataset.transform.e == pytest.approx(-250.0)
            assert dataset.transform.c == pytest.approx(427875.0)
            assert dataset.transform.f == pytest.approx(3365625.0)
            assert dataset.bounds.left == pytest.approx(427875.0)
            assert dataset.bounds.right == pytest.approx(428625.0)
            assert dataset.bounds.bottom == pytest.approx(3364875.0)
            assert dataset.bounds.top == pytest.approx(3365625.0)
            assert values[0, 2] == pytest.approx(99.0)
            assert values[2, 1] == pytest.approx(GEOTIFF_NODATA)


def test_geotiff_custom_epsg_and_local_unknown_crs():
    rasterio = pytest.importorskip("rasterio")
    grid_x, grid_y, grid_z = _asymmetric_grid()

    custom = grid_to_geotiff_bytes(grid_x, grid_y, grid_z, {"mode": "Custom EPSG", "epsg": 3857})
    local = grid_to_geotiff_bytes(grid_x, grid_y, grid_z, local_crs().to_dict())

    with rasterio.io.MemoryFile(custom) as memory_file:
        with memory_file.open() as dataset:
            assert dataset.crs.to_epsg() == 3857
    with rasterio.io.MemoryFile(local) as memory_file:
        with memory_file.open() as dataset:
            assert dataset.crs is None
            assert dataset.bounds.left == pytest.approx(427875.0)


def test_zmap_round_trip_header_orientation_null_and_crs_comments():
    grid_x, grid_y, grid_z = _asymmetric_grid()
    metadata = build_map_export_metadata(_map_result(), nodata=ZMAP_NULL_VALUE)
    zmap = grid_to_zmap_ascii_bytes(grid_x, grid_y, grid_z, metadata=metadata)
    text = zmap.decode("utf-8")
    parsed = parse_zmap_grid_ascii_bytes(zmap)

    assert "@GRID FILE, GRID, 4" in text
    assert "! CRS: EPSG:32638" in text
    assert "! VALUE_ORDER: north-to-south rows, west-to-east columns." in text
    assert parsed["nx"] == 3
    assert parsed["ny"] == 3
    assert parsed["x_min"] == pytest.approx(428000.0)
    assert parsed["x_max"] == pytest.approx(428500.0)
    assert parsed["y_min"] == pytest.approx(3365000.0)
    assert parsed["y_max"] == pytest.approx(3365500.0)
    assert parsed["dx"] == pytest.approx(250.0)
    assert parsed["dy"] == pytest.approx(250.0)
    assert parsed["null_value"] == pytest.approx(ZMAP_NULL_VALUE)
    assert parsed["grid_z"][2, 2] == pytest.approx(99.0)
    assert np.isnan(parsed["grid_z"][0, 1])


def test_kriging_geotiff_uncertainty_files_share_grid_and_orientation():
    rasterio = pytest.importorskip("rasterio")
    files = map_geotiff_export_files(_map_result(variance=True), project_metadata={"name": "Rumaila"})
    names = [name for name, _, _ in files]

    assert any(name.endswith("_Estimate.tif") for name in names)
    assert any(name.endswith("_KrigingVariance.tif") for name in names)
    assert any(name.endswith("_KrigingStdDev.tif") for name in names)

    transforms = []
    crs_codes = []
    north_east_values = {}
    for name, data, _ in files:
        with rasterio.io.MemoryFile(data) as memory_file:
            with memory_file.open() as dataset:
                transforms.append(dataset.transform)
                crs_codes.append(dataset.crs.to_epsg())
                north_east_values[name] = dataset.read(1)[0, 2]
    assert len(set(transforms)) == 1
    assert set(crs_codes) == {32638}
    assert next(value for name, value in north_east_values.items() if "Estimate" in name) == pytest.approx(99.0)
    assert next(value for name, value in north_east_values.items() if "KrigingStdDev" in name) == pytest.approx(9.0)


def test_engineering_control_exports_and_excel_sheets_preserve_traceability():
    from openpyxl import load_workbook

    map_result = _map_result()
    controls = engineering_controls_export_dataframe(map_result)
    regions = control_regions_export_dataframe(map_result)
    metadata = build_map_export_metadata(map_result)
    workbook = load_workbook(
        BytesIO(
            grid_to_excel_bytes(
                pd.DataFrame({"X": [1.0], "Y": [2.0], "Value": [3.0]}),
                metadata,
                engineering_controls=controls,
                control_regions=regions,
                validation_df=pd.DataFrame([{"RMSE": 1.5}]),
            )
        )
    )

    assert list(controls.columns) == [
        "Control_ID",
        "Source_Type",
        "Region_ID",
        "X",
        "Y",
        "Property",
        "Value",
        "Unit",
        "Reservoir_Layer",
        "Panel",
        "Pressure_Reference_Date",
        "Active",
        "Comment",
    ]
    assert controls.loc[1, "Source_Type"] == REGION_CONTROL_SOURCE
    assert controls.loc[1, "Region_ID"] == "CR-001"
    assert regions.loc[0, "Region_ID"] == "CR-001"
    assert metadata["Control_Regions"][0]["Geometry"]["type"] == "Polygon"
    assert {"Interpolated_Grid", "Map_Metadata", "Engineering_Controls", "Control_Regions", "Validation"}.issubset(
        set(workbook.sheetnames)
    )


def test_batch_all_layer_package_exports_one_file_per_layer_and_batch_metadata():
    upper = _map_result(layer="Upper", variance=False)
    lower = _map_result(layer="Lower", variance=False)
    package = batch_map_package_zip_bytes(
        {"Upper": upper, "Lower": lower},
        project_metadata={"name": "Rumaila"},
        include_formats={"geotiff", "zmap", "metadata"},
    )

    with zipfile.ZipFile(BytesIO(package)) as archive:
        names = archive.namelist()
        batch_metadata = json.loads(archive.read("metadata.json").decode("utf-8"))

    assert sum(name.endswith(".tif") for name in names) == 2
    assert sum(name.endswith(".zmap") for name in names) == 2
    assert any("Upper" in name and name.endswith(".tif") for name in names)
    assert any("Lower" in name and name.endswith(".zmap") for name in names)
    assert batch_metadata["Generated_Reservoir_Layers"] == ["Lower", "Upper"]
    assert batch_metadata["CRS"]["epsg"] == 32638


def test_saved_scenario_export_uses_snapshot_crs_after_workspace_changes():
    rasterio = pytest.importorskip("rasterio")
    generated = _map_result(crs=preset_crs_32638().to_dict())
    scenario = create_map_scenario(
        "Pressure Upper 2026",
        generated,
        {"crs": local_crs().to_dict()},
    )
    generated["crs"] = local_crs().to_dict()
    opened = scenario_to_generated_map(scenario)
    metadata = build_map_export_metadata(opened, scenario=scenario)
    files = map_geotiff_export_files(opened, scenario=scenario)

    assert metadata["Scenario_Name"] == "Pressure Upper 2026"
    assert metadata["CRS_EPSG"] == 32638
    with rasterio.io.MemoryFile(files[0][1]) as memory_file:
        with memory_file.open() as dataset:
            assert dataset.crs.to_epsg() == 32638
            assert dataset.read(1)[0, 2] == pytest.approx(99.0)


def test_project_crs_and_scenario_crs_round_trip():
    generated = _map_result(crs=preset_crs_32638().to_dict())
    scenario = create_map_scenario("Pressure Upper 2026", generated)
    state = {
        "project_metadata": {"name": "CRS Round Trip"},
        "source_name": "sample.csv",
        "column_mappings": {"x": "X", "y": "Y", "property": "Pressure"},
        "coordinate_unit": "meters",
        "property_unit": "psi",
        "crs": preset_crs_32638().to_dict(),
        "geometry_layers": {"reservoir_boundary": None, "panels": None, "faults": None, "custom": []},
        "map_scenarios": [scenario],
        "working_df": pd.DataFrame({"X": [428000.0], "Y": [3365000.0], "Pressure": [3000.0]}),
    }

    restored = load_project_archive(save_project_archive(state))
    restored_scenario = restored["map_scenarios"][0]

    assert restored["crs"]["mode"] == EPSG_32638_CRS_MODE
    assert restored["crs"]["epsg"] == 32638
    assert restored["crs"]["name"] == "WGS 84 / UTM zone 38N"
    assert restored["coordinate_unit"] == "meters"
    assert restored_scenario["crs"]["epsg"] == 32638
    assert restored_scenario["crs"]["name"] == "WGS 84 / UTM zone 38N"


def test_zmap_pressure_change_metadata_records_later_minus_earlier():
    grid_x, grid_y, grid_z = _asymmetric_grid()
    delta_map = {
        "grid_x": grid_x,
        "grid_y": grid_y,
        "grid_z": grid_z,
        "property_col": "Pressure Change",
        "property_type": "Pressure Change",
        "unit": "psi",
        "coordinate_unit": "meters",
        "crs": preset_crs_32638().to_dict(),
        "method": "Delta",
        "metadata": {
            "Operation": "Later - Earlier",
            "Earlier_Reference_Date": "2025-01-01",
            "Later_Reference_Date": "2026-01-01",
            "Negative_Values": "Pressure decline",
            "Positive_Values": "Pressure increase",
        },
    }

    zmap_name, zmap_data, metadata_name, metadata_data, metadata = map_zmap_export_files(delta_map)[0]

    assert zmap_name.startswith("PressureChange_2025-01-01_to_2026-01-01")
    assert metadata_name.endswith("_metadata.json")
    assert b"Pressure Change" in zmap_data
    parsed_metadata = json.loads(metadata_data.decode("utf-8"))
    assert parsed_metadata["Operation"] == "Later - Earlier"
    assert parsed_metadata["Negative_Values"] == "Pressure decline"
    assert metadata["Later_Reference_Date"] == "2026-01-01"
